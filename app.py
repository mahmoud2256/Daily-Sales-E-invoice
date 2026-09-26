import io, re, xml.etree.ElementTree as ET
from datetime import datetime
import pandas as pd
import streamlit as st
import openpyxl

st.set_page_config(page_title="Daily Sales E-invoice", page_icon="🧾", layout="wide")

CSS = """
<style>
.stApp { background-color: #0F172A; color: #E2E8F0; }
header[data-testid="stHeader"] { background: #0F172A; }
section[data-testid="stSidebar"]{display:none;}
.header-card{background:linear-gradient(135deg,#1E293B 0%,#0F172A 100%);border:1px solid #334155;border-radius:16px;padding:16px 20px;display:flex;justify-content:space-between;align-items:center;margin-bottom:16px;}
.card{background:#1E293B;border:1px solid #334155;border-radius:14px;padding:12px;}
.badge{font-size:10px;padding:2px 6px;border-radius:8px;margin-top:4px;display:inline-block;}
.badge-req{background:#1E3A8A;color:#93C5FD;border:1px solid #3B82F6;}
.badge-tpl{background:#0F766E;color:#99F6E4;border:1px solid #14B8A6;}
.badge-val{background:#14532D;color:#86EFAC;border:1px solid #22C55E;}
[data-testid="stFileUploaderDropzone"]{background-color:#0F172A !important;border:1px dashed #475569 !important;}
div.stButton > button{background:#2563EB !important;color:white !important;font-weight:700 !important;border:none !important;border-radius:10px !important;}
.stDownloadButton > button{background:#059669 !important;color:white !important;font-weight:700 !important;border:none !important;border-radius:10px !important;}
.metric-card{background:#1E293B;border:1px solid #334155;border-radius:12px;padding:12px;}
</style>
<div class="header-card">
  <div style="display:flex;align-items:center;gap:10px;"><div style="width:44px;height:44px;background:#2563EB;border-radius:10px;display:flex;align-items:center;justify-content:center;">🧾</div><div><div style="font-size:18px;font-weight:800;color:#F8FAFC;">Daily Sales E-invoice</div><div style="font-size:11px;color:#94A3B8;">Full Dashboard - Cloud Fixed</div></div></div>
  <div style="color:#22C55E;font-size:11px;">✅ App Loaded - openpyxl OK</div>
</div>
"""
st.markdown(CSS, unsafe_allow_html=True)

d1,d2,d3,d4=st.columns(4)
d1.markdown('<div class="metric-card"><div style="color:#94A3B8;font-size:11px;">Tax</div><div style="font-size:13px;font-weight:800;color:#22C55E;">Auto-Detect 14 ✓</div></div>', unsafe_allow_html=True)
d2.markdown('<div class="metric-card"><div style="color:#94A3B8;font-size:11px;">Date</div><div style="font-size:13px;font-weight:700;color:#3B82F6;">YYYY-MM-DD Only</div></div>', unsafe_allow_html=True)
d3.markdown('<div class="metric-card"><div style="color:#94A3B8;font-size:11px;">USD</div><div style="font-size:13px;font-weight:700;">amtegp + rate</div></div>', unsafe_allow_html=True)
d4.markdown('<div class="metric-card"><div style="color:#94A3B8;font-size:11px;">PO/Requester</div><div style="font-size:13px;font-weight:700;">From All Processed</div></div>', unsafe_allow_html=True)

def parse_xml_spreadsheet(uploaded):
    try:
        uploaded.seek(0)
        content=uploaded.read()
        if isinstance(content, bytes):
            content=content.decode('utf-8', errors='ignore')
        import tempfile
        with tempfile.NamedTemporaryFile(mode='w', suffix='.xml', delete=False, encoding='utf-8') as tmp:
            tmp.write(content)
            tmp_path=tmp.name
        ns={'ss':'urn:schemas-microsoft-com:office:spreadsheet'}
        tree=ET.parse(tmp_path)
        ws=tree.getroot().find('.//ss:Worksheet', ns)
        rows=ws.findall('.//ss:Row', ns)
        header_cells=rows[0].findall('.//ss:Cell', ns)
        headers=[c.find('ss:Data', ns).text if c.find('ss:Data', ns) is not None else "" for c in header_cells]
        data=[]
        for r in rows[1:]:
            cells=r.findall('.//ss:Cell', ns)
            row_vals=[]; col_idx=0
            for cell in cells:
                idx=cell.attrib.get('{urn:schemas-microsoft-com:office:spreadsheet}Index')
                if idx:
                    target=int(idx)-1
                    while col_idx<target:
                        row_vals.append(""); col_idx+=1
                d=cell.find('ss:Data', ns)
                row_vals.append(d.text if d is not None else ""); col_idx+=1
            while len(row_vals)<len(headers):
                row_vals.append("")
            data.append(row_vals)
        return pd.DataFrame(data, columns=headers)
    except:
        return None

def read_file(uploaded):
    if not uploaded:
        return None
    try:
        uploaded.seek(0)
        name=uploaded.name.lower()
        if 'all' in name and 'processed' in name:
            df=parse_xml_spreadsheet(uploaded)
            if df is not None and len(df)>0:
                return df
        if name.endswith('.csv'):
            uploaded.seek(0)
            return pd.read_csv(uploaded)
        if name.endswith('.xls') and not name.endswith('.xlsx'):
            df=parse_xml_spreadsheet(uploaded)
            if df is not None and len(df)>0:
                return df
            try:
                uploaded.seek(0)
                return pd.read_excel(uploaded, engine='xlrd')
            except:
                uploaded.seek(0)
                return pd.read_excel(uploaded)
        uploaded.seek(0)
        try:
            return pd.read_excel(uploaded, engine='openpyxl')
        except:
            uploaded.seek(0)
            return pd.read_excel(uploaded)
    except:
        return None

def clean(v):
    if pd.isna(v):
        return ""
    s=str(v).strip()
    return "" if s.lower() in ['nan','none','nat'] else re.sub(r'\s+',' ', s)

def get_receiver_type(tax_id):
    s=re.sub(r'\D','',str(tax_id).replace('.0','').strip())
    return ("أجنبي","تصدير للخارج") if len(s)>9 and s else ("شركة","سلع عامة")

def find_col(df,kws):
    for kw in kws:
        for col in df.columns:
            if kw.lower() in str(col).lower():
                return col
    return None

def get_val(row,names,default=""):
    for name in names:
        if name in row.index and pd.notna(row[name]) and str(row[name]).strip()!="":
            return row[name]
    for col in row.index:
        cl=str(col).lower()
        for name in names:
            if name.lower() in cl and pd.notna(row[col]) and str(row[col]).strip()!="":
                return row[col]
    return default

def detect_tax_columns(df):
    candidates=[]
    for col in df.columns:
        try:
            vals=pd.to_numeric(df[col], errors='coerce')
            count_14=((vals==14) | (vals==14.0)).sum()
            count_0=(vals==0).sum()
            total=vals.notna().sum()
            if count_14>0 and total>0:
                if count_14 + count_0 >= total*0.7:
                    candidates.append((col,count_14))
        except:
            continue
    candidates.sort(key=lambda x: x[1], reverse=True)
    if candidates:
        return candidates[0][0]
    for kw in ['taxter','tax rate','النسبة','tax_rate','inv_tax']:
        c=find_col(df,[kw])
        if c:
            return c
    return None

c1,c2,c3=st.columns(3)
with c1:
    st.markdown('<div class="card"><h4>📊 Sales Data (EG)</h4><span class="badge badge-req">Required</span></div>', unsafe_allow_html=True)
    sales_file=st.file_uploader("sales",type=["xlsx","xls","csv"],key="f1",label_visibility="collapsed")
with c2:
    st.markdown('<div class="card"><h4>📋 All Processed</h4><span class="badge badge-req">PO & Requester</span></div>', unsafe_allow_html=True)
    all_proc_file=st.file_uploader("all",type=["xlsx","xls","csv"],key="f2",label_visibility="collapsed")
with c3:
    st.markdown('<div class="card"><h4>📄 Empty Template</h4><span class="badge badge-tpl">Template</span></div>', unsafe_allow_html=True)
    template_file=st.file_uploader("template",type=["xlsx"],key="f4",label_visibility="collapsed")

st.markdown("---")
d1,d2,d3=st.columns([2,3,1])
with d1:
    inv_date=st.date_input("Invoice Date",value=datetime(2026,9,23))
    inv_date_only=inv_date.strftime("%Y-%m-%d")
with d2:
    st.markdown("<div style='height:24px'></div>", unsafe_allow_html=True)
    run_btn=st.button("⚡ Generate",use_container_width=True,type="primary")
with d3:
    st.markdown("<div style='height:24px'></div>", unsafe_allow_html=True)
    if st.button("Reset",use_container_width=True):
        st.rerun()

if run_btn:
    if not sales_file or not template_file:
        st.error("Upload Sales EG + Template")
        st.stop()
    sales_df=read_file(sales_file)
    all_proc_df=read_file(all_proc_file) if all_proc_file else None
    if sales_df is None or len(sales_df)==0:
        st.error("Sales file empty")
        st.stop()
    tax_col_auto=detect_tax_columns(sales_df)
    if tax_col_auto:
        st.info(f"Detected Tax Column: {tax_col_auto} | 14 count: {(pd.to_numeric(sales_df[tax_col_auto], errors='coerce')==14).sum()}")
    lookup={}
    if all_proc_df is not None:
        client_col=find_col(all_proc_df,['Client Name'])
        req_col=find_col(all_proc_df,['Requester'])
        po_col=find_col(all_proc_df,['PO NUM'])
        for _, r in all_proc_df.iterrows():
            client=clean(r[client_col]) if client_col else ""
            req=clean(r[req_col]) if req_col else ""
            po=clean(r[po_col]) if po_col else ""
            if client and (req or po):
                if client not in lookup:
                    lookup[client]={'req':req,'po':po}
                else:
                    if req and not lookup[client]['req']:
                        lookup[client]['req']=req
                    if po and not lookup[client]['po']:
                        lookup[client]['po']=po
    rows=[]
    for _, r in sales_df.iterrows():
        internalid=clean(get_val(r,['internalid','مسلسل الفاتورة','Invoice'], ""))
        if not internalid or len(internalid)<2 or 'internalid' in internalid.lower():
            continue
        tax_id_raw=clean(get_val(r,['rec_id','رقم ت.ض','Company ID'], ""))
        receiver_type,_=get_receiver_type(tax_id_raw)
        curr=clean(get_val(r,['inv_unitval_currsold','العملة'], "EGP"))
        amount_col_egp=find_col(sales_df,['inv_unitva_inv_unitva','inv_salestotal','السعر قبل الخصم'])
        amount_col_usd=find_col(sales_df,['inv_unitval_amtegp','amtegp'])
        rate_col=find_col(sales_df,['inv_unitval_currrate','currrate','سعر العملة'])
        amount_val=0; exchange_rate=""
        if curr.upper()=="USD":
            if amount_col_usd:
                try: amount_val=float(get_val(r,[amount_col_usd],0))
                except: amount_val=0
            if rate_col: exchange_rate=clean(get_val(r,[rate_col],""))
        else:
            if amount_col_egp:
                try: amount_val=float(get_val(r,[amount_col_egp],0))
                except: amount_val=0
            if amount_val==0:
                try: amount_val=float(get_val(r,['inv_salestotal','inv_nettotal'],0))
                except: amount_val=0
        tax_rate_val=""; tax_class=""; tax_sub=""
        tax_col_to_use=tax_col_auto if tax_col_auto else find_col(sales_df,['inv_taxiter_inv_taxter','taxter','النسبة'])
        if tax_col_to_use:
            raw=get_val(r,[tax_col_to_use],"")
            try:
                tr=float(raw)
                if tr!=0 and str(raw).strip() not in ["","0","0.0"]:
                    tax_rate_val=int(tr) if tr==int(tr) else tr
                    tax_class="ضريبه القيمه المضافه"
                    tax_subtype_col=find_col(sales_df,['inv_taxiter_inv_taxiter','taxiter','التصنيف الفرعى'])
                    if tax_subtype_col:
                        st_raw=clean(get_val(r,[tax_subtype_col],""))
                        tax_sub=st_raw if st_raw and st_raw!="0" else "سلع عامة"
                    else:
                        tax_sub="سلع عامة" if receiver_type=="شركة" else "تصدير للخارج"
            except:
                if str(raw).strip() not in ["","0"]:
                    tax_rate_val=raw
                    tax_class="ضريبه القيمه المضافه"
                    tax_sub="سلع عامة"
        client_name=clean(get_val(r,['rec_name','الاسم','Invoicing name'], ""))
        po_num=lookup.get(client_name,{}).get('po',"") if client_name in lookup else ""
        requester=lookup.get(client_name,{}).get('req',"") if client_name in lookup else ""
        if not po_num: po_num=clean(get_val(r,['PO NUM','مرجع طلب الشراء'],""))
        if not requester: requester=clean(get_val(r,['Requester','مرجع أمر البيع'],""))
        row_out={
            'مسلسل الفاتورة':internalid,'تاريخ الاصدار':inv_date_only,'نوع الوثيقة':"فاتورة",
            'مرجع طلب الشراء':po_num,'مرجع أمر البيع':requester,'رقم الفاتورة المبدئية للتصدير':'',
            'المرجع':clean(get_val(r,['References','المرجع'],"")),'قيمة الخصم الاضافى':'','الصفة':receiver_type,
            'الاسم':client_name,'رقم ت.ض - رقم قومى':tax_id_raw,'الدولة':clean(get_val(r,['rec_add_governate.1','الدولة'],"مصر")),
            'المحافظة':clean(get_val(r,['Area','المحافظة'],"Cairo")),'الحى':clean(get_val(r,['rec_add_regioncity.1','الحى'],"")),
            'الشارع':clean(get_val(r,['rec_add_street.1','الشارع'],"")),'رقم المبنى':clean(get_val(r,['rec_add_buildingnumber.1','رقم المبنى'],""))[:60],
            'الوصف':clean(get_val(r,['inv_description','الوصف'],"")),'كود المنتج (EGS-GS1)':clean(get_val(r,['inv_itemcode','كود المنتج'],"")),
            'الكود الداخلى':clean(get_val(r,['inv_itemcode.1','الكود الداخلى'],"")),'وحدة القياس':clean(get_val(r,['inv_unitty','وحدة القياس'],"EA")),
            'الكمية':get_val(r,['inv_quantity','الكمية'],1),'السعر قبل الخصم':amount_val,'قيمة الخصم':'','فرق القيمة':'',
            'العملة':curr if curr else "EGP",'معامل تحويل العملة':exchange_rate if curr.upper()=="USD" else "",
            'الضريبة 1 - التصنيف':tax_class,'الضريبة 1 - التصنيف الفرعى':tax_sub,'الضريبة 1 - النسبة':tax_rate_val,
            'الضريبة 2 - التصنيف':'','الضريبة 2 - التصنيف الفرعى':'','الضريبة 2 - النسبة':'',
            'الضريبة 3 - التصنيف':'','الضريبة 3 - التصنيف الفرعى':'','الضريبة 3 - النسبة ':'',
            'الضريبة 4 - التصنيف':'','الضريبة 4 - التصنيف الفرعى':'','الضريبة 4 - النسبة ':'',
        }
        rows.append(row_out)
    invoices_df=pd.DataFrame(rows)
    st.metric("Total", len(invoices_df))
    st.dataframe(invoices_df[['مسلسل الفاتورة','الاسم','السعر قبل الخصم','العملة','الضريبة 1 - النسبة']].head(20), use_container_width=True)
    output=io.BytesIO()
    template_file.seek(0)
    wb=openpyxl.load_workbook(template_file)
    ws=wb['Invoices']
    max_row=ws.max_row
    if max_row>2: ws.delete_rows(3, max_row-2)
    headers=[ws.cell(row=2,column=c).value for c in range(1, ws.max_column+1)]
    for r_idx, row_data in invoices_df.iterrows():
        excel_row=3+r_idx
        for c_idx, header in enumerate(headers, start=1):
            if header in row_data:
                cell=ws.cell(row=excel_row,column=c_idx,value=row_data[header])
                if 'تاريخ الاصدار' in str(header): cell.number_format='YYYY-MM-DD'
            else:
                for k in row_data.keys():
                    if str(k).strip()==str(header).strip():
                        cell=ws.cell(row=excel_row,column=c_idx,value=row_data[k])
                        if 'تاريخ الاصدار' in str(header): cell.number_format='YYYY-MM-DD'
                        break
    for tbl in ws._tables:
        try: tbl.ref=f"A2:AL{len(invoices_df)+2}"
        except: pass
    wb.save(output)
    st.download_button("⬇️ Download", data=output.getvalue(), file_name=f"Daily_Sales_{inv_date_only}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True)
